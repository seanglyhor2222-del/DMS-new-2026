from datetime import datetime, date, time
from types import SimpleNamespace
import io
import json
import os

from flask import Blueprint, render_template, request, redirect, url_for, flash, jsonify, send_file, current_app
from flask_login import login_required, current_user

from extensions import db
from models import DocumentRecord, LookupItem, SystemSetting, AuditLog
from routes.auth import permission_required

documents_bp = Blueprint("documents", __name__)

SORTABLE_COLUMNS = {
    "seq_no": DocumentRecord.seq_no,
    "doc_type": DocumentRecord.doc_type,          
    "doc_number": DocumentRecord.doc_number,
    "doc_date": DocumentRecord.doc_date,
    "action_date": DocumentRecord.action_date,  # ✅ ប្រើ action_date ឱ្យស៊ីចង្វាក់គ្នា
    "priority": DocumentRecord.priority,
    "status": DocumentRecord.status,
}

EXPORT_COLUMNS = [
    ("ល.រ", "seq_no"),
    ("ចេញ/ចូល", "doc_type"),
    ("លេខលិខិត", "doc_number"),
    ("កាលបរិច្ឆេទលិខិត", "doc_date"),
    ("កម្មវត្ថុ", "subject"),
    ("អង្គភាព", "unit"),
    ("ចំនួន", "quantity"),
    ("អាទិភាព", "priority"),
    ("កាលបរិច្ឆេទ ចូល", "in_date"),
    ("កាលបរិច្ឆេទ ចេញ", "out_date"),
    ("ម៉ោង ចូល", "in_time"),
    ("ម៉ោង ចេញ", "out_time"),
    ("ស្ថានភាព", "status"),
    ("កំណត់សម្គាល់", "remarks"),
    ("ផ្សេងៗ", "others"),
]


def _audit_snapshot(record):
    return record.to_dict()


def _audit_detail(action, before=None, after=None):
    return json.dumps({
        "format": "document_snapshot_v1",
        "action": action,
        "before": before,
        "after": after,
    }, ensure_ascii=False)


def _next_seq_no():
    last = DocumentRecord.query.order_by(DocumentRecord.seq_no.desc()).first()
    return (last.seq_no + 1) if last else 1


def _next_doc_number():
    n = int(SystemSetting.get("doc_number_next", "1"))
    SystemSetting.set("doc_number_next", n + 1)
    return str(n)


def _is_duplicate_doc_number(doc_number, exclude_id=None):
    if not doc_number:
        return False
    q = DocumentRecord.query.filter(DocumentRecord.doc_number == doc_number)
    if exclude_id is not None:
        q = q.filter(DocumentRecord.id != exclude_id)
    return db.session.query(q.exists()).scalar()


def _get_action_fields(f):
    in_date = _parse_date(f.get("in_date"))
    out_date = _parse_date(f.get("out_date"))
    in_time = _parse_time(f.get("in_time"))
    out_time = _parse_time(f.get("out_time"))
    doc_type = f.get("doc_type", "").strip()
    action_date = in_date or out_date
    action_time = in_time or out_time
    return in_date, out_date, in_time, out_time, action_date, action_time


def _pseudo_record_from_form(f):
    in_date, out_date, in_time, out_time, action_date, action_time = _get_action_fields(f)
    return SimpleNamespace(
        id=None,
        seq_no=_parse_int(f.get("seq_no")),
        doc_type=f.get("doc_type", "").strip() or None,
        doc_number=f.get("doc_number", ""),
        doc_date=_parse_date(f.get("doc_date")) or date.today(),
        subject=f.get("subject", ""),
        unit=f.get("unit", ""),
        quantity=_parse_int(f.get("quantity")),
        priority=f.get("priority", ""),
        action_date=action_date,
        action_time=action_time,
        in_date=in_date,
        out_date=out_date,
        in_time=in_time,
        out_time=out_time,
        status=f.get("status", ""),
        remarks=f.get("remarks", ""),
        others=f.get("others", ""),
    )


def _build_filtered_query(args):
    q = args.get("q", "").strip()
    date_from = args.get("date_from", "")
    date_to = args.get("date_to", "")
    doc_type = args.get("doc_type", "")
    sort_by = args.get("sort_by", "seq_no")
    sort_dir = args.get("sort_dir", "desc")

    query = DocumentRecord.query

    def parse_filter_date(value):
        for date_format in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(value, date_format).date()
            except ValueError:
                continue
        return None

    if q and current_user.permission("can_search"):
        like = f"%{q}%"
        query = query.filter(
            db.or_(DocumentRecord.doc_number.ilike(like),
                   DocumentRecord.subject.ilike(like),
                   DocumentRecord.unit.ilike(like))
        )

    if date_from:
        df = parse_filter_date(date_from)
        if df:
            query = query.filter(DocumentRecord.in_date == df)

    if date_to:
        dt = parse_filter_date(date_to)
        if dt:
            query = query.filter(DocumentRecord.out_date == dt)

    if doc_type == "__empty__":
        query = query.filter(db.or_(DocumentRecord.doc_type.is_(None),
                                    DocumentRecord.doc_type == "",
                                    DocumentRecord.doc_type == "-"))
    elif doc_type:
        query = query.filter(DocumentRecord.doc_type == doc_type)

    col = SORTABLE_COLUMNS.get(sort_by, DocumentRecord.seq_no)
    query = query.order_by(col.desc() if sort_dir == "desc" else col.asc())
    parsed_date_from = parse_filter_date(date_from) if date_from else None
    parsed_date_to = parse_filter_date(date_to) if date_to else None
    date_from_display = parsed_date_from.strftime("%d/%m/%Y") if parsed_date_from else date_from
    date_to_display = parsed_date_to.strftime("%d/%m/%Y") if parsed_date_to else date_to

    return query, {"q": q, "date_from": date_from, "date_to": date_to,
                   "date_from_display": date_from_display, "date_to_display": date_to_display,
                   "doc_type": doc_type,
                   "sort_by": sort_by, "sort_dir": sort_dir}


@documents_bp.route("/")
@login_required
def dashboard():
    query, ctx = _build_filtered_query(request.args)
    records = query.all()

    lookups = {
        "doc_type": LookupItem.query.filter_by(category="doc_type", is_active=True).order_by(LookupItem.sort_order).all(),
        "unit": LookupItem.query.filter_by(category="unit", is_active=True).order_by(LookupItem.sort_order).all(),
        "priority": LookupItem.query.filter_by(category="priority", is_active=True).order_by(LookupItem.sort_order).all(),
        "status": LookupItem.query.filter_by(category="status", is_active=True).order_by(LookupItem.sort_order).all(),
    }

    settings = {
        "doc_number_mode": SystemSetting.get("doc_number_mode", "auto"),
        "doc_type_free_text": SystemSetting.get("doc_type_free_text", "true") == "true",
        "unit_free_text": SystemSetting.get("unit_free_text", "true") == "true",
        "priority_free_text": SystemSetting.get("priority_free_text", "true") == "true",
        "status_free_text": SystemSetting.get("status_free_text", "true") == "true",
        "freeze_table_header": SystemSetting.get("freeze_table_header") != "false",
        "date_input_mode": SystemSetting.get("date_input_mode", "auto"),
    }

    # Build the filter from both saved records and active lookup values.
    available_doc_types = []
    seen_doc_types = set()

    raw_doc_types = db.session.query(DocumentRecord.doc_type).distinct().all()
    for (value,) in raw_doc_types:
        value = value.strip() if isinstance(value, str) else value
        if value and value != "-" and value not in seen_doc_types:
            available_doc_types.append(value)
            seen_doc_types.add(value)

    for item in lookups["doc_type"]:
        value = item.value_km.strip() if isinstance(item.value_km, str) else item.value_km
        if value and value != "-" and value not in seen_doc_types:
            available_doc_types.append(value)
            seen_doc_types.add(value)

    available_doc_types.sort(key=str.casefold)
    available_doc_types.insert(0, "__empty__")

    return render_template("dashboard.html", records=records, lookups=lookups, settings=settings, 
                           available_doc_types=available_doc_types, **ctx)


@documents_bp.route("/documents/export/xlsx")
@login_required
def export_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    query, ctx = _build_filtered_query(request.args)
    records = query.all()

    wb = Workbook()
    ws = wb.active
    ws.title = "ឯកសារ"

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for col_idx, (label, _) in enumerate(EXPORT_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[chr(64 + col_idx)].width = max(12, len(label) + 4)

    for row_idx, r in enumerate(records, start=2):
        for col_idx, (_, attr) in enumerate(EXPORT_COLUMNS, start=1):
            val = getattr(r, attr)
            if attr == "doc_type":
                val = val if val else ""
            elif attr in ("doc_date", "action_date") and val:
                val = val.strftime("%d-%m-%Y")
            elif attr == "action_time" and val:
                val = val.strftime("%H:%M")
            ws.cell(row=row_idx, column=col_idx, value=val)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    db.session.add(AuditLog(user_id=current_user.id, action="export_xlsx",
                            detail=f"Exported {len(records)} records to Excel"))
    db.session.commit()

    filename = f"documents_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return send_file(buf, as_attachment=True, download_name=filename,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@documents_bp.route("/documents/export/pdf")
@login_required
def export_pdf():
    query, ctx = _build_filtered_query(request.args)
    records = query.all()

    db.session.add(AuditLog(user_id=current_user.id, action="export_pdf",
                            detail=f"Exported {len(records)} records to PDF"))
    db.session.commit()

    return render_template("print_documents.html", records=records,
                           generated_at=datetime.now(), filters=ctx)


@documents_bp.route("/documents/import/template.xlsx")
@login_required
@permission_required("can_add")
def import_template_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "ទម្រង់នាំចូល"

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    col_idx = 0
    for label, attr in EXPORT_COLUMNS:
        if attr == "seq_no":
            continue
        col_idx += 1
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = header_font
        cell.fill = header_fill
        ws.column_dimensions[chr(64 + col_idx)].width = max(14, len(label) + 4)

    example = ["ចូល", "001", "01-01-2026", "កម្មវត្ថុគំរូ", "នាយកដ្ឋានរដ្ឋបាល", 1,
               "ធម្មតា", "01-01-2026", "09:00", "រួចរាល់", "", ""]
    for col_idx, val in enumerate(example, start=1):
        ws.cell(row=2, column=col_idx, value=val)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="ទម្រង់_នាំចូល.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@documents_bp.route("/documents/import", methods=["GET", "POST"])
@login_required
@permission_required("can_add")
def import_documents():
    if request.method == "GET":
        return render_template("doc_import.html", imported=None, skipped=None)

    file = request.files.get("xlsx_file")
    if not file or file.filename == "":
        flash("សូមជ្រើសរើសឯកសារ Excel (.xlsx) ។", "danger")
        return redirect(url_for("documents.import_documents"))

    from openpyxl import load_workbook
    try:
        wb = load_workbook(file, data_only=True)
    except Exception:
        flash("មិនអាចអានឯកសារនេះបានទេ។ សូមប្រាកដថាវាជាឯកសារ .xlsx ត្រឹមត្រូវ។", "danger")
        return redirect(url_for("documents.import_documents"))

    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        flash("ឯកសារនេះមិនមានទិន្នន័យទេ។", "danger")
        return redirect(url_for("documents.import_documents"))

    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    col_index = {}
    for label, attr in EXPORT_COLUMNS:
        if attr == "seq_no":
            continue
        if label in header:
            col_index[attr] = header.index(label)

    if "doc_number" not in col_index and "subject" not in col_index:
        flash("ក្បាលតារាងមិនត្រូវនឹងទម្រង់នាំចូលទេ។ សូមទាញយក \"ទម្រង់គំរូ\" ខាងក្រោមជាមុនសិន រួចបំពេញឡើងវិញ។", "danger")
        return redirect(url_for("documents.import_documents"))

    def cell(row, attr):
        idx = col_index.get(attr)
        if idx is None or idx >= len(row):
            return None
        return row[idx]

    imported = 0
    skipped = []
    for i, row in enumerate(rows[1:], start=2):
        doc_number = str(cell(row, "doc_number") or "").strip()
        subject = str(cell(row, "subject") or "").strip()
        if not doc_number and not subject:
            continue

        if doc_number and _is_duplicate_doc_number(doc_number):
            skipped.append(f"ជួរទី {i}: លេខលិខិត \"{doc_number}\" មានរួចហើយក្នុងប្រព័ន្ធ")
            continue

        doc_type = str(cell(row, "doc_type") or "").strip() or None

        rec = DocumentRecord(
            seq_no=_next_seq_no(),
            doc_type=doc_type,
            doc_number=doc_number,
            doc_number_is_auto=False,
            doc_date=_parse_excel_date(cell(row, "doc_date")) or date.today(),
            subject=subject,
            unit=str(cell(row, "unit") or "").strip(),
            quantity=_parse_int(cell(row, "quantity")),
            priority=str(cell(row, "priority") or "").strip(),
            in_date=_parse_excel_date(cell(row, "in_date")),
            out_date=_parse_excel_date(cell(row, "out_date")),
            in_time=_parse_excel_time(cell(row, "in_time")),
            out_time=_parse_excel_time(cell(row, "out_time")),
            status=str(cell(row, "status") or "").strip(),
            remarks=str(cell(row, "remarks") or "").strip(),
            others=str(cell(row, "others") or "").strip(),
            created_by_id=current_user.id,
        )
        db.session.add(rec)
        imported += 1

    if imported:
        db.session.commit()
        db.session.add(AuditLog(
            user_id=current_user.id, action="import_xlsx",
            detail=f"Imported {imported} records from Excel" + (f", skipped {len(skipped)} duplicates" if skipped else "")
        ))
        db.session.commit()
        flash(f"បាននាំចូល {imported} កំណត់ត្រាដោយជោគជ័យ។" +
              (f" (រំលង {len(skipped)} ជួរ ព្រោះស្ទួនលេខលិខិត)" if skipped else ""), "success")
    else:
        flash("មិនមានកំណត់ត្រាណាមួយត្រូវបាននាំចូលទេ (ប្រហែលជាទិន្នន័យទទេ ឬស្ទួនទាំងអស់)។", "danger")

    return render_template("doc_import.html", imported=imported, skipped=skipped)


@documents_bp.route("/documents/new", methods=["GET", "POST"])
@login_required
@permission_required("can_add")
def create_document():
    lookups = {
        "doc_type": LookupItem.query.filter_by(category="doc_type", is_active=True).order_by(LookupItem.sort_order).all(),
        "unit": LookupItem.query.filter_by(category="unit", is_active=True).order_by(LookupItem.sort_order).all(),
        "priority": LookupItem.query.filter_by(category="priority", is_active=True).order_by(LookupItem.sort_order).all(),
        "status": LookupItem.query.filter_by(category="status", is_active=True).order_by(LookupItem.sort_order).all(),
    }
    settings = {
        "doc_number_mode": SystemSetting.get("doc_number_mode", "auto"),
        "doc_type_free_text": SystemSetting.get("doc_type_free_text", "true") == "true",
        "unit_free_text": SystemSetting.get("unit_free_text", "true") == "true",
        "priority_free_text": SystemSetting.get("priority_free_text", "true") == "true",
        "status_free_text": SystemSetting.get("status_free_text", "true") == "true",
        "date_input_mode": SystemSetting.get("date_input_mode", "auto"),
    }

    if request.method == "POST":
        f = request.form

        doc_number_mode = settings["doc_number_mode"]
        if doc_number_mode == "auto":
            doc_number = _next_doc_number()
            is_auto = True
        else:
            doc_number = f.get("doc_number", "").strip()
            is_auto = False

        if _is_duplicate_doc_number(doc_number):
            return render_template(
                "doc_form.html", record=_pseudo_record_from_form(f), lookups=lookups, settings=settings,
                mode="create", today=date.today().isoformat(), now_time=datetime.now().strftime("%H:%M"),
                doc_number_error=f"លេខលិខិត \"{doc_number}\" មានរួចហើយ! សូមប្រើលេខផ្សេង។",
            )

        in_date, out_date, in_time, out_time, action_date, action_time = _get_action_fields(f)

        rec = DocumentRecord(
            seq_no=_next_seq_no(),
            doc_type=f.get("doc_type", "").strip() or None,
            doc_number=doc_number,
            doc_number_is_auto=is_auto,
            doc_date=_parse_date(f.get("doc_date")) or date.today(),
            subject=f.get("subject", "").strip(),
            unit=f.get("unit", "").strip(),
            quantity=_parse_int(f.get("quantity")),
            priority=f.get("priority", "").strip(),
            action_date=action_date,
            action_time=action_time,
            in_date=in_date,
            out_date=out_date,
            in_time=in_time,
            out_time=out_time,
            status=f.get("status", "").strip(),
            remarks=f.get("remarks", "").strip(),
            others=f.get("others", "").strip(),
            created_by_id=current_user.id,
        )
        db.session.add(rec)
        db.session.commit()
        db.session.add(AuditLog(user_id=current_user.id, action="create",
                                record_id=rec.id,
                                detail=_audit_detail("create", after=_audit_snapshot(rec))))
        db.session.commit()
        flash("បានរក្សាទុកឯកសារដោយជោគជ័យ។", "success")
        return redirect(url_for("documents.dashboard"))

    return render_template("doc_form.html", record=None, lookups=lookups, settings=settings, mode="create",
                           today=date.today().isoformat(), now_time=datetime.now().strftime("%H:%M"))


@documents_bp.route("/documents/<int:doc_id>/edit", methods=["GET", "POST"])
@login_required
@permission_required("can_edit")
def edit_document(doc_id):
    rec = DocumentRecord.query.get_or_404(doc_id)
    lookups = {
        "doc_type": LookupItem.query.filter_by(category="doc_type", is_active=True).order_by(LookupItem.sort_order).all(),
        "unit": LookupItem.query.filter_by(category="unit", is_active=True).order_by(LookupItem.sort_order).all(),
        "priority": LookupItem.query.filter_by(category="priority", is_active=True).order_by(LookupItem.sort_order).all(),
        "status": LookupItem.query.filter_by(category="status", is_active=True).order_by(LookupItem.sort_order).all(),
    }
    settings = {
        "doc_number_mode": SystemSetting.get("doc_number_mode", "auto"),
        "doc_type_free_text": SystemSetting.get("doc_type_free_text", "true") == "true",
        "unit_free_text": SystemSetting.get("unit_free_text", "true") == "true",
        "priority_free_text": SystemSetting.get("priority_free_text", "true") == "true",
        "status_free_text": SystemSetting.get("status_free_text", "true") == "true",
        "date_input_mode": SystemSetting.get("date_input_mode", "auto"),
    }

    if request.method == "POST":
        f = request.form
        before_snapshot = _audit_snapshot(rec)

        if current_user.is_admin() and f.get("seq_no"):
            rec.seq_no = _parse_int(f.get("seq_no")) or rec.seq_no

        new_doc_number = rec.doc_number
        if current_user.is_admin():
            new_doc_number = f.get("doc_number", rec.doc_number)
        elif settings["doc_number_mode"] == "manual":
            new_doc_number = f.get("doc_number", rec.doc_number)

        if new_doc_number != rec.doc_number and _is_duplicate_doc_number(new_doc_number, exclude_id=rec.id):
            pseudo = _pseudo_record_from_form(f)
            pseudo.id = rec.id
            pseudo.seq_no = rec.seq_no if pseudo.seq_no is None else pseudo.seq_no
            return render_template(
                "doc_form.html", record=pseudo, lookups=lookups, settings=settings, mode="edit",
                today=date.today().isoformat(), now_time=datetime.now().strftime("%H:%M"),
                doc_number_error=f"លេខលិខិត \"{new_doc_number}\" មានរួចហើយ! សូមប្រើលេខផ្សេង។",
            )
        rec.doc_number = new_doc_number

        in_date, out_date, in_time, out_time, action_date, action_time = _get_action_fields(f)

        rec.doc_type = f.get("doc_type", rec.doc_type).strip() or None
        rec.doc_date = _parse_date(f.get("doc_date")) or rec.doc_date
        rec.subject = f.get("subject", rec.subject)
        rec.unit = f.get("unit", rec.unit)
        rec.quantity = _parse_int(f.get("quantity"))
        rec.priority = f.get("priority", rec.priority)
        rec.action_date = action_date
        rec.action_time = action_time
        rec.in_date = in_date
        rec.out_date = out_date
        rec.in_time = in_time
        rec.out_time = out_time
        rec.status = f.get("status", rec.status)
        rec.remarks = f.get("remarks", rec.remarks)
        rec.others = f.get("others", rec.others)
        rec.updated_by_id = current_user.id
        rec.updated_at = datetime.utcnow()

        db.session.add(AuditLog(user_id=current_user.id, action="edit",
                                record_id=rec.id,
                                detail=_audit_detail("edit", before=before_snapshot,
                                                     after=_audit_snapshot(rec))))
        db.session.commit()
        flash("បានកែប្រែឯកសារដោយជោគជ័យ។", "success")
        return redirect(url_for("documents.dashboard"))

    return render_template("doc_form.html", record=rec, lookups=lookups, settings=settings, mode="edit",
                           today=date.today().isoformat(), now_time=datetime.now().strftime("%H:%M"))


@documents_bp.route("/documents/<int:doc_id>/delete", methods=["POST"])
@login_required
def delete_document(doc_id):
    if not current_user.is_admin():
        flash("មានតែ Admin ទេអាចលុបបាន។", "danger")
        return redirect(url_for("documents.dashboard"))
    rec = DocumentRecord.query.get_or_404(doc_id)
    db.session.add(AuditLog(user_id=current_user.id, action="delete",
                            record_id=rec.id,
                            detail=_audit_detail("delete", before=_audit_snapshot(rec))))
    db.session.delete(rec)
    db.session.commit()
    flash("បានលុបឯកសារ។", "success")
    return redirect(url_for("documents.dashboard"))


@documents_bp.route("/documents/bulk-delete", methods=["POST"])
@login_required
def bulk_delete_documents():
    if not current_user.is_admin():
        flash("មានតែ Admin ទេអាចលុបបាន។", "danger")
        return redirect(url_for("documents.dashboard"))

    doc_ids = request.form.getlist("document_ids")
    if not doc_ids:
        flash("សូមជ្រើសរើសលិខិតយ៉ាងហោចណាស់មួយដើម្បីលុប!", "warning")
        return redirect(url_for("documents.dashboard"))

    deleted_count = 0
    for doc_id_str in doc_ids:
        try:
            doc_id = int(doc_id_str)
            rec = DocumentRecord.query.get(doc_id)
            if rec:
                db.session.add(AuditLog(
                    user_id=current_user.id, 
                    action="delete",
                    record_id=rec.id, 
                    detail=_audit_detail("delete", before=_audit_snapshot(rec))
                ))
                db.session.delete(rec)
                deleted_count += 1
        except ValueError:
            continue

    if deleted_count > 0:
        db.session.commit()
        flash(f"បានលុបឯកសារចំនួន {deleted_count} ដោយជោគជ័យ។", "success")
    else:
        flash("មិនមានឯកសារណាមួយត្រូវបានលុបទេ។", "danger")

    return redirect(url_for("documents.dashboard"))


def _parse_date(s):
    if not s:
        return None
    try:
        return datetime.strptime(str(s).strip(), "%Y-%m-%d").date()
    except ValueError:
        return None


def _parse_time(s):
    if not s:
        return None
    try:
        return datetime.strptime(str(s).strip(), "%H:%M").time()
    except ValueError:
        return None


def _parse_excel_date(val):
    if val is None or val == "":
        return None
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    s = str(val).strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _parse_excel_time(val):
    if val is None or val == "":
        return None
    if isinstance(val, datetime):
        return val.time()
    if isinstance(val, time):
        return val
    s = str(val).strip()
    for fmt in ("%H:%M", "%H:%M:%S"):
        try:
            return datetime.strptime(s, fmt).time()
        except ValueError:
            continue
    return None


def _parse_int(s):
    try:
        return int(s)
    except (TypeError, ValueError):
        return None